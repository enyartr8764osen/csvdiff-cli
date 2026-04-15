"""Render diff results as an Excel (.xlsx) workbook."""
from __future__ import annotations

from typing import List

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "openpyxl is required for Excel output: pip install openpyxl"
    ) from exc

from csvdiff.core import DiffResult

# Colour palette (ARGB hex strings)
_GREEN = "FF92D050"
_RED = "FFFF4C4C"
_YELLOW = "FFFFC000"
_HEADER = "FF4472C4"


def _header_row(ws, columns: List[str]) -> None:
    """Write a bold, coloured header row to *ws*."""
    fill = PatternFill(fill_type="solid", fgColor=_HEADER)
    font = Font(bold=True, color="FFFFFFFF")
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = fill
        cell.font = font


def _write_section(
    ws,
    rows: list,
    columns: List[str],
    start_row: int,
    fgColor: str,
) -> int:
    """Append *rows* to *ws* starting at *start_row*; return next free row."""
    fill = PatternFill(fill_type="solid", fgColor=fgColor)
    for row in rows:
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=row.get(col, ""))
            cell.fill = fill
        start_row += 1
    return start_row


def render_excel(diff: DiffResult, path: str) -> None:
    """Write *diff* to an Excel workbook at *path*.

    Three sheets are produced:
      - Added   – rows present only in the new file (green)
      - Removed – rows present only in the old file (red)
      - Changed – old/new pairs for modified rows (yellow)
    """
    wb = openpyxl.Workbook()
    columns = diff.columns

    for sheet_name, rows, colour in (
        ("Added", diff.added, _GREEN),
        ("Removed", diff.removed, _RED),
    ):
        ws = wb.create_sheet(title=sheet_name)
        _header_row(ws, columns)
        _write_section(ws, rows, columns, start_row=2, fgColor=colour)

    ws_changed = wb.create_sheet(title="Changed")
    header = [f"{c} (old)" for c in columns] + [f"{c} (new)" for c in columns]
    _header_row(ws_changed, header)
    row_idx = 2
    fill = PatternFill(fill_type="solid", fgColor=_YELLOW)
    for old_row, new_row in diff.changed:
        for col_idx, col in enumerate(columns, start=1):
            cell = ws_changed.cell(row=row_idx, column=col_idx, value=old_row.get(col, ""))
            cell.fill = fill
        for col_idx, col in enumerate(columns, start=1):
            cell = ws_changed.cell(row=row_idx, column=len(columns) + col_idx, value=new_row.get(col, ""))
            cell.fill = fill
        row_idx += 1

    # Remove the default empty sheet created by openpyxl
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(path)
